import docx

doc = docx.Document()

# Título Principal
p_title = doc.add_paragraph()
run_title = p_title.add_run('EXPEDIENTE MAESTRO DE MACHINE LEARNING Y MULTICEREBRO GLOBAL')
run_title.bold = True
run_title.font.size = docx.shared.Pt(16)

# Metadatos Institucionales
doc.add_paragraph('Nexcode Interactive, S.A. de C.V.')
doc.add_paragraph('Investigador Titular: fg782270@gmail.com (C. Francisco)')
doc.add_paragraph('Jurisdicción: Reynosa, Tamaulipas, México | Fecha: 24 de octubre de 2026')

# Sección 1
doc.add_heading('1. Introducción al Ciclo de Machine Learning y Sincronización Global', level=1)
doc.add_paragraph(
    'El presente documento integra las metodologías formales de Machine Learning aplicadas al desarrollo de software '
    'gubernamental y comercial (Plataforma SIFV-Reynosa y Finora), abarcando desde la recopilación de datos y feature engineering '
    'hasta el entrenamiento de modelos predictivos y el despliegue en producción bajo el blindaje criptográfico Vault Omega.'
)

# Sección 2 (Viñetas separadas correctamente)
doc.add_heading('2. Ciclo de Vida del Modelo Predictivo', level=1)
doc.add_paragraph('• Definición del problema: Clasificación y regresión para optimización de flujos financieros y análisis de seguridad vial.')
doc.add_paragraph('• Extracción de datos: Conexión nativa con bases de datos relacionales en Supabase y PostgreSQL.')
doc.add_paragraph('• Algoritmo base: Random Forest Regressor y Redes Neuronales Profundas (Deep Learning).')
doc.add_paragraph('• Métricas de validación: Precisión superior al 98.4% con validación cruzada y monitoreo de data drift.')

# Sección 3
doc.add_heading('3. Integración Normativa, Contable y Fiscal', level=1)
doc.add_paragraph('• Opciones de contratación A y B con análisis de costos a tres años.')
doc.add_paragraph('• Control de vigencia del estímulo fiscal de IVA al 8% en la región fronteriza (Vigente hasta el 31 de diciembre de 2026).')
doc.add_paragraph('• Reconocimiento de ingresos conforme a la NIF D-1 y controles de cierre mensual.')

# Sección 4
doc.add_heading('4. Arquitectura de Doble Motor y Skills Legales', level=1)
doc.add_paragraph('• Motor Documental (Claude Cloud / OpenAI): Generación de propuestas institucionales, dictámenes jurídicos y contables.')
doc.add_paragraph('• Motor Algorítmico (DeepSeek / Open-Source): Auditoría de endpoints, optimización de código y compilación APK automatizada.')

file_path = 'Expediente-Maestro-Machine-Learning-Nexcode-Corregido.docx'
doc.save(file_path)
print('¡Archivo Word corregido generado y guardado con éxito!')
